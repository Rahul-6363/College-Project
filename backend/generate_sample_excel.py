"""
Generate a sample block-format attendance Excel file for testing.
Run: python generate_sample_excel.py
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import date, timedelta
import random

EMPLOYEES = [
    ("EMP001", "Alice Johnson", "Engineering"),
    ("EMP002", "Bob Smith", "HR"),
    ("EMP003", "Carol White", "Finance"),
    ("EMP004", "David Brown", "Engineering"),
    ("EMP005", "Eva Martinez", "Operations"),
]

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
CUSTOM_HOLIDAYS = [7, 21]

def gen_time(base_hour, base_min, variance=30):
    mins = base_min + random.randint(-variance, variance)
    h = base_hour + mins // 60
    m = mins % 60
    return f"{h:02d}:{m:02d}"

def create_sample_excel(filename="sample_attendance.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance April 2025"

    row = 1
    start_date = date(2025, 4, 1)

    for empcode, name, dept in EMPLOYEES:
        # Employee header row
        ws.cell(row=row, column=1, value="Emp. Code")
        ws.cell(row=row, column=2, value=empcode)
        ws.cell(row=row, column=3, value="Employee Name")
        ws.cell(row=row, column=4, value=name)
        for col in range(1, 5):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="2D2D3A")
        row += 1

        # Department row
        ws.cell(row=row, column=1, value="Dept. Name")
        ws.cell(row=row, column=2, value=dept)
        ws.cell(row=row, column=1).font = Font(bold=True)
        row += 1

        # Column headers
        headers = ["Date", "Day", "In Time", "Out Time", "Work Hour", "Status"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="3A3A52")
        row += 1

        # Data rows for 30 days
        for d in range(30):
            current = start_date + timedelta(days=d)

            # Skip Sundays
            if current.weekday() == 6:
                continue
            # Skip custom holidays
            if current.day in CUSTOM_HOLIDAYS:
                continue

            day_abbr = DAYS[current.weekday()]
            date_str = current.strftime("%d-%b-%Y")

            # Random attendance
            rand = random.random()
            if rand < 0.85:  # 85% present
                status = "P"
                in_time = gen_time(9, 0, 40)  # Around 9am ± 40min
                out_time = gen_time(17, 30, 30)

                # Compute work hours simply
                in_h, in_m = map(int, in_time.split(":"))
                out_h, out_m = map(int, out_time.split(":"))
                work_mins = (out_h * 60 + out_m) - (in_h * 60 + in_m)
                work_hour = f"{work_mins // 60}:{work_mins % 60:02d}"
            else:
                status = "A"
                in_time = "----"
                out_time = "----"
                work_hour = "----"

            ws.cell(row=row, column=1, value=date_str)
            ws.cell(row=row, column=2, value=day_abbr)
            ws.cell(row=row, column=3, value=in_time)
            ws.cell(row=row, column=4, value=out_time)
            ws.cell(row=row, column=5, value=work_hour)
            ws.cell(row=row, column=6, value=status)
            row += 1

        # Blank separator
        row += 2

    # Auto-size columns
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 20)

    wb.save(filename)
    print(f"✅ Sample Excel file created: {filename}")
    print(f"   Employees: {len(EMPLOYEES)}")
    print(f"   Upload this to the Data Manager dashboard to test.")

if __name__ == "__main__":
    create_sample_excel()